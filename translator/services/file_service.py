import io
import csv
import json
import logging

import markdown as md_lib
from PyPDF2 import PdfReader
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation as PptxPresentation

from .openai_service import translate_document_text

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {
    ".txt", ".csv", ".json", ".xml", ".html", ".htm",
    ".md", ".markdown", ".srt",
    ".pdf", ".docx", ".xlsx", ".pptx",
}


def get_file_extension(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


async def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract text content from various file types."""
    ext = get_file_extension(filename)

    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes)
    elif ext == ".xlsx":
        return _extract_xlsx(file_bytes)
    elif ext == ".pptx":
        return _extract_pptx(file_bytes)
    elif ext == ".csv":
        return _extract_csv(file_bytes)
    elif ext == ".json":
        return _extract_json(file_bytes)
    elif ext in (".xml", ".html", ".htm"):
        return file_bytes.decode("utf-8", errors="replace")
    elif ext in (".md", ".markdown"):
        return file_bytes.decode("utf-8", errors="replace")
    elif ext == ".srt":
        return file_bytes.decode("utf-8", errors="replace")
    elif ext == ".txt":
        return file_bytes.decode("utf-8", errors="replace")
    else:
        return file_bytes.decode("utf-8", errors="replace")


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n\n".join(text_parts)


def _extract_docx(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    parts = []
    for para in doc.paragraphs:
        parts.append(para.text)
    return "\n".join(parts)


def _extract_pptx(file_bytes: bytes) -> str:
    prs = PptxPresentation(io.BytesIO(file_bytes))
    text_parts = []
    for slide in prs.slides:
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    slide_texts.append(para.text)
        text_parts.append("\n".join(slide_texts))
    return "\n\n---\n\n".join(text_parts)


def _extract_xlsx(file_bytes: bytes) -> str:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    text_parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            rows.append("\t".join(cells))
        text_parts.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    return "\n\n".join(text_parts)


def _extract_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    return text


def _extract_json(file_bytes: bytes) -> str:
    text = file_bytes.decode("utf-8", errors="replace")
    try:
        data = json.loads(text)
        return json.dumps(data, indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return text


async def translate_file(
    file_bytes: bytes,
    filename: str,
    target_language: str,
    source_language: str = "auto",
) -> tuple[bytes, str, str]:
    """
    Translate a file. Returns (translated_bytes, output_filename, content_type).
    """
    ext = get_file_extension(filename)
    text = await extract_text_from_file(file_bytes, filename)

    translated_text = await translate_document_text(
        text, target_language, source_language
    )

    output_filename = _make_output_filename(filename, target_language)

    if ext == ".docx":
        output_bytes = _create_docx(translated_text)
        content_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    elif ext == ".xlsx":
        output_bytes = _create_xlsx(translated_text)
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    elif ext == ".json":
        output_bytes = translated_text.encode("utf-8")
        content_type = "application/json"
    elif ext == ".csv":
        output_bytes = translated_text.encode("utf-8")
        content_type = "text/csv"
    elif ext in (".html", ".htm"):
        output_bytes = translated_text.encode("utf-8")
        content_type = "text/html"
    elif ext in (".xml",):
        output_bytes = translated_text.encode("utf-8")
        content_type = "application/xml"
    elif ext in (".md", ".markdown"):
        output_bytes = translated_text.encode("utf-8")
        content_type = "text/markdown"
    elif ext == ".srt":
        output_bytes = translated_text.encode("utf-8")
        content_type = "application/x-subrip"
    else:
        output_bytes = translated_text.encode("utf-8")
        content_type = "text/plain"

    return output_bytes, output_filename, content_type


def _make_output_filename(filename: str, target_language: str) -> str:
    if "." in filename:
        name, ext = filename.rsplit(".", 1)
        return f"{name}_translated_{target_language}.{ext}"
    return f"{filename}_translated_{target_language}"


def _create_docx(text: str) -> bytes:
    doc = DocxDocument()
    for para_text in text.split("\n"):
        doc.add_paragraph(para_text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _create_xlsx(text: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Translated"
    for i, line in enumerate(text.split("\n"), 1):
        cells = line.split("\t")
        for j, cell in enumerate(cells, 1):
            ws.cell(row=i, column=j, value=cell)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
